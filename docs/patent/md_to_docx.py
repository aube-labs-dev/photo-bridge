#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Photo Bridge 명세서 MD → KIPO 표준 docx (HI_FIVE.docx 서식 동일)."""
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import re, os

BASE = "/Users/jonghwa/dev/source/photo-bridge/docs/patent"
MD = f"{BASE}/PHOTO_BRIDGE_PATENT_SPECIFICATION_KR.md"
FIG_DIR = f"{BASE}/patent_figures"
OUT = "/Users/jonghwa/Documents/PHOTO_BRIDGE.docx"

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Batang'
style.font.size = Pt(11)
rf = style.element.get_or_add_rPr().get_or_add_rFonts()
rf.set(qn('w:eastAsia'), '바탕')
rf.set(qn('w:ascii'), 'Batang')
rf.set(qn('w:hAnsi'), 'Batang')
style.paragraph_format.line_spacing = 1.6
for sec in doc.sections:
    sec.top_margin = Mm(25); sec.bottom_margin = Mm(20)
    sec.left_margin = Mm(20); sec.right_margin = Mm(20)

TAG = re.compile(r'^【.+】$')

def set_ea(run):
    r = run.element.get_or_add_rPr().get_or_add_rFonts()
    r.set(qn('w:eastAsia'), '바탕')

def add_tag(t):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run(t); r.bold = True; set_ea(r)

def add_body(t):
    p = doc.add_paragraph()
    r = p.add_run(t); set_ea(r)

with open(MD, encoding='utf-8') as f:
    for raw in f:
        line = raw.rstrip('\n')
        s = line.strip()
        if not s:
            continue
        if TAG.match(s):
            add_tag(s)
        else:
            add_body(s)

# 도면 섹션
add_tag('【도면】')
for i, fn in enumerate(['fig1.png', 'fig2.png', 'fig3.png', 'fig4.png', 'fig5.png'], 1):
    add_tag(f'【도 {i}】')
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(os.path.join(FIG_DIR, fn), width=Mm(150))

doc.save(OUT)
print("saved:", OUT)
